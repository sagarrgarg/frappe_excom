"""What a person did, and who is allowed to ask.

The permission half matters more than the figures. A report is exactly the kind of screen somebody
tries by changing a name in the URL, so "an agent sees only themselves" has to hold in the endpoint
and not merely in the page that calls it — and it has to hold for the download too, which is where
it would otherwise be forgotten.
"""

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import add_days, add_to_date, now_datetime, today

from excom.excom.api import reports
from excom.excom.tests.fixtures import purge_user

AGENT = "qa.report.agent@example.com"
MATE = "qa.report.mate@example.com"
BOSS = "qa.report.boss@example.com"
USERS = (AGENT, MATE, BOSS)


def _cleanup():
	frappe.set_user("Administrator")
	for oi in frappe.get_all("Omni Identity", {"display_name": ["like", "QA Report%"]}, pluck="name"):
		for t in frappe.get_all("Excom Thread", {"omni_identity": oi}, pluck="name"):
			frappe.db.delete("Excom Message", {"thread": t})
			frappe.delete_doc("Excom Thread", t, force=True, ignore_permissions=True)
		for child in ("Omni Identity Link", "Omni Identity Channel", "Omni Identity Alias"):
			frappe.db.delete(child, {"parent": oi})
		frappe.delete_doc("Omni Identity", oi, force=True, ignore_permissions=True)
	for t in frappe.get_all("ToDo", {"description": ["like", "QA Report%"]}, pluck="name"):
		frappe.delete_doc("ToDo", t, force=True, ignore_permissions=True)
	for u in USERS:
		purge_user(u)
	frappe.db.commit()


class TestActivityReport(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		_cleanup()
		for email, role in ((AGENT, "Excom Agent"), (MATE, "Excom Agent"), (BOSS, "Excom Admin")):
			u = frappe.get_doc({
				"doctype": "User", "email": email, "first_name": email.split("@")[0],
				"send_welcome_email": 0,
			})
			u.flags.ignore_permissions = True
			u.insert(ignore_permissions=True)
			u.add_roles(role)

		ref = frappe.get_all(
			"Excom Thread", fields=["channel", "account_doctype", "account"], limit=1
		)[0]
		identity = frappe.get_doc({
			"doctype": "Omni Identity", "display_name": "QA Report Buyer",
			"primary_phone": "+919900000771",
		}).insert(ignore_permissions=True).name
		cls.thread = frappe.get_doc({
			"doctype": "Excom Thread", "omni_identity": identity, "channel": ref.channel,
			"account_doctype": ref.account_doctype, "account": ref.account,
			"thread_key": "qa-report-1", "status": "Open",
			"last_message_at": now_datetime(),
		}).insert(ignore_permissions=True).name
		cls.ref = ref
		cls.identity = identity
		frappe.db.commit()

	@classmethod
	def tearDownClass(cls):
		_cleanup()
		super().tearDownClass()

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.db.delete("Excom Message", {"thread": self.thread})
		for t in frappe.get_all("ToDo", {"description": ["like", "QA Report%"]}, pluck="name"):
			frappe.delete_doc("ToDo", t, force=True, ignore_permissions=True)
		frappe.db.commit()

	def _message(self, by, when=None, channel=None):
		frappe.set_user("Administrator")
		doc = frappe.get_doc({
			"doctype": "Excom Message", "thread": self.thread,
			"omni_identity": self.identity, "direction": "Outbound",
			"channel": channel or self.ref.channel,
			"account_doctype": self.ref.account_doctype, "account": self.ref.account,
			"message_type": "Text", "content_text": "QA Report hello",
			"provider_message_id": frappe.generate_hash(length=12),
		})
		doc.flags.ignore_permissions = True
		doc.insert(ignore_permissions=True)
		frappe.db.set_value(
			"Excom Message", doc.name,
			{"created_by_user": by, "creation": when or now_datetime()},
			update_modified=False,
		)
		frappe.db.commit()
		return doc.name

	# ── who may ask ───────────────────────────────────────────────────────────

	def test_an_agent_sees_only_themselves(self):
		frappe.set_user(AGENT)
		r = reports.get_activity_report()
		self.assertEqual([row["user"] for row in r["rows"]], [AGENT])

	def test_an_agent_cannot_ask_about_a_colleague(self):
		"""The report is where somebody would try another name in the URL."""
		frappe.set_user(AGENT)
		with self.assertRaises(frappe.PermissionError):
			reports.get_activity_report(user=MATE)

	def test_an_admin_sees_the_whole_desk(self):
		frappe.set_user(BOSS)
		r = reports.get_activity_report()
		users = {row["user"] for row in r["rows"]}
		self.assertIn(AGENT, users)
		self.assertIn(MATE, users)

	def test_an_admin_may_ask_about_one_person(self):
		frappe.set_user(BOSS)
		r = reports.get_activity_report(user=AGENT)
		self.assertEqual([row["user"] for row in r["rows"]], [AGENT])

	def test_the_download_obeys_the_same_rule(self):
		"""A file is where the permission check is most easily left out."""
		from excom.excom.api import report_files

		frappe.set_user(AGENT)
		with self.assertRaises(frappe.PermissionError):
			report_files.download_activity_report(user=MATE, fmt="xlsx")

	def test_reportable_users_is_one_person_for_an_agent(self):
		frappe.set_user(AGENT)
		self.assertEqual([u["name"] for u in reports.reportable_users()], [AGENT])

	# ── the window ────────────────────────────────────────────────────────────

	def test_a_day_is_that_day(self):
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="daily", on=today())
		self.assertEqual(r["from"], r["to"])

	def test_a_week_runs_monday_to_sunday(self):
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="weekly", on="2026-09-16")  # a Wednesday
		self.assertEqual(r["from"], "2026-09-14")
		self.assertEqual(r["to"], "2026-09-20")

	def test_a_month_is_the_whole_month(self):
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="monthly", on="2026-09-16")
		self.assertEqual(r["from"], "2026-09-01")
		self.assertEqual(r["to"], "2026-09-30")

	def test_an_unknown_period_is_refused(self):
		frappe.set_user(BOSS)
		with self.assertRaises(frappe.ValidationError):
			reports.get_activity_report(period="fortnightly")

	# ── the figures ───────────────────────────────────────────────────────────

	def _row(self, report, user):
		return next(r for r in report["rows"] if r["user"] == user)

	def test_messages_are_counted_against_whoever_sent_them(self):
		self._message(AGENT)
		self._message(AGENT)
		self._message(MATE)
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="daily", on=today())
		self.assertEqual(self._row(r, AGENT)["messages"]["total"], 2)
		self.assertEqual(self._row(r, MATE)["messages"]["total"], 1)

	def test_work_outside_the_window_is_not_counted(self):
		self._message(AGENT, when=add_to_date(now_datetime(), days=-10))
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="daily", on=today())
		self.assertEqual(self._row(r, AGENT)["messages"]["total"], 0)

	def test_a_week_includes_a_day_inside_it(self):
		self._message(AGENT, when=add_to_date(now_datetime(), days=-1))
		frappe.set_user(BOSS)
		daily = reports.get_activity_report(period="daily", on=today())
		weekly = reports.get_activity_report(period="weekly", on=today())
		self.assertEqual(self._row(daily, AGENT)["messages"]["total"], 0)
		self.assertGreaterEqual(self._row(weekly, AGENT)["messages"]["total"], 1)

	def test_tasks_made_and_finished_are_both_counted(self):
		"""Two different things a person did, so a task made and finished the same day counts once
		in each rather than cancelling out."""
		frappe.set_user("Administrator")
		doc = frappe.get_doc({
			"doctype": "ToDo", "description": "QA Report task", "date": today(),
			"allocated_to": AGENT, "status": "Closed",
		})
		doc.flags.ignore_permissions = True
		doc.insert(ignore_permissions=True)
		frappe.db.set_value("ToDo", doc.name, "owner", AGENT, update_modified=False)
		frappe.db.commit()

		frappe.set_user(BOSS)
		row = self._row(reports.get_activity_report(period="daily", on=today()), AGENT)
		self.assertEqual(row["tasks"]["created"], 1)
		self.assertEqual(row["tasks"]["completed"], 1)

	def test_somebody_who_did_nothing_still_appears_and_says_so(self):
		"""A quiet week is a fact about the week, not a person missing from the report."""
		frappe.set_user(BOSS)
		row = self._row(reports.get_activity_report(period="daily", on=today()), MATE)
		self.assertFalse(row["did_anything"])
		self.assertEqual(row["messages"]["total"], 0)

	def test_the_totals_add_up(self):
		self._message(AGENT)
		self._message(MATE)
		frappe.set_user(BOSS)
		r = reports.get_activity_report(period="daily", on=today())
		self.assertEqual(
			r["totals"]["messages"]["total"],
			sum(row["messages"]["total"] for row in r["rows"]),
		)


class TestTheFiles(FrappeTestCase):
	"""A report nobody can send on is half a report."""

	def setUp(self):
		frappe.set_user("Administrator")

	def test_the_workbook_is_a_real_xlsx(self):
		from excom.excom.api import report_files

		report = reports.get_activity_report(period="weekly")
		data = report_files._xlsx(report)
		self.assertTrue(data.startswith(b"PK"), "an xlsx is a zip, and this one is not")

		import io as _io
		from openpyxl import load_workbook

		wb = load_workbook(_io.BytesIO(data))
		self.assertIn("Activity", wb.sheetnames)
		self.assertIn("Breakdown", wb.sheetnames)
		ws = wb["Activity"]
		self.assertEqual(ws.cell(row=5, column=1).value, "Person")
		self.assertEqual(
			ws.cell(row=5, column=1).fill.fgColor.rgb[-6:], report_files.ACCENT,
			"the header should carry the accent, or the file is not the coloured one",
		)

	def test_the_html_carries_every_person_and_a_total(self):
		from excom.excom.api import report_files

		report = reports.get_activity_report(period="weekly")
		html = report_files._html(report)
		for row in report["rows"]:
			self.assertIn(row["full_name"], html)
		self.assertIn("All %d" % report["totals"]["people"], html)

	def test_an_unknown_format_is_refused(self):
		from excom.excom.api import report_files

		with self.assertRaises(frappe.ValidationError):
			report_files.download_activity_report(fmt="docx")
